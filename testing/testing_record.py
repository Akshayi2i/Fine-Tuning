"""The testing record: one row per document an extraction run read, in the
columns of the team's workbook ("Testing Record" sheet of
``Documentation/Testing Records/FIdeon SLM - Model Training & Test Record_v1.xlsx``).

The extraction command writes the rows as it goes (``testing/metrics/<model>/
test_record.csv``, :func:`record_row`); this command copies them into the
workbook, under the round they belong to::

    python -m testing.testing_record --model base --round Baseline \\
        --xlsx "Documentation/Testing Records/FIdeon SLM - Model Training & Test Record_v1.xlsx"

**With no project code on the pod** (it serves only ``vllm serve`` and the
MinerU API), ``testing.latency_check`` runs here and reads the documents through
them. ``--stage`` lays the chosen bundles out as it reads them - each PDF and
its gold under one short document id, the bundle's own name kept in a local
index - and ``--latency-dir`` turns its output into the same rows::

    python -m testing.testing_record --stage --bundles data/bundles_c1 --select baseline \\
        --out testing/results/baseline_c1                      # prints the --doc arguments
    python -m testing.latency_check --model-url https://<pod>-8000.proxy.runpod.net \\
        --mineru-url https://<pod>-8001.proxy.runpod.net --mode image_only \\
        --out testing/results/baseline_c1/run "@testing/results/baseline_c1/docs.txt"
    python -m testing.testing_record --latency-dir testing/results/baseline_c1/run \\
        --index testing/results/baseline_c1/doc_index.csv --round Baseline --xlsx "<workbook>"

What each column measures:

* **F1 Score Accuracy** - the document's field match rate on its scalar fields
  (``field_normalized_match_rate``, the gate's matching rules). Each scalar
  field has one predicted and one expected value, so precision, recall and F1
  coincide (pilot/zero_shot_baseline.py); table rows are scored apart, as list
  recall, in the run's metrics. Empty when the gold was not scored.
* **Latency** - seconds the model spent generating the document's answer
  (``ExtractionResult.latency_ms``; ``latency_check``'s ``model`` timing); OCR
  and page loading are not in it.
* **Input / Output Tokens** - summed over every window read, retries included
  (``window_stats``): what the document cost the GPU.
* **Carrier** - as the gold names it (``carrier.name``), else as the label
  metadata records it.
"""

from __future__ import annotations

import argparse
import csv
from collections.abc import Iterable, Sequence
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Any

#: The workbook's sheet and its headers, in order (matched without surrounding spaces).
SHEET = "Testing Record"
COLUMNS = ("Training Round", "Training/Testing Date", "Test/Val", "Carrier", "LOB", "Doc ID", "# Pages",
           "F1 Score Accuracy", "Latency", "Input Tokens", "Output Tokens")
#: Beside the sheet's columns in the CSV: what read the document, and how.
EXTRA = ("Model", "Mode")
RECORD_FILE = "test_record.csv"


class RecordError(RuntimeError):
    pass


def _carrier(golden: dict[str, Any] | None, metadata: dict[str, Any]) -> str | None:
    from common.canonical import values_view

    carrier = values_view(golden or {}).get("carrier")
    if isinstance(carrier, dict):
        # `name` on the common model, `company_name` on a self-contained schema.
        carrier = carrier.get("name") or carrier.get("company_name")
    if isinstance(carrier, str) and carrier.strip():
        return carrier.strip()
    return metadata.get("carrier") or None


def _tokens(result: Any, key: str) -> int | None:
    counts = [stat.get(key) for stat in getattr(result, "window_stats", None) or []]
    counts = [count for count in counts if isinstance(count, int)]
    return sum(counts) if counts else None


def record_row(result: Any, metrics: dict[str, Any], *, metadata: dict[str, Any], golden: dict[str, Any] | None,
               pages: int, round_label: str, date: str | None = None) -> dict[str, Any]:
    """One document's row: the sheet's columns, then :data:`EXTRA`."""
    match = metrics.get("field_normalized_match_rate")
    return {
        "Training Round": round_label,
        "Training/Testing Date": date or datetime.now(UTC).date().isoformat(),
        "Test/Val": metadata.get("split"),
        "Carrier": _carrier(golden, metadata),
        "LOB": metadata.get("lob"),
        "Doc ID": result.source_id,
        "# Pages": pages,
        "F1 Score Accuracy": round(match, 4) if isinstance(match, (int, float)) else None,
        "Latency": round(result.latency_ms / 1000, 1) if result.latency_ms is not None else None,
        "Input Tokens": _tokens(result, "prompt_tokens"),
        "Output Tokens": _tokens(result, "answer_tokens"),
        "Model": result.model_version,
        "Mode": result.mode,
    }


def append_csv(rows: Iterable[dict[str, Any]], path: Path) -> None:
    """Add rows to the run's record, writing the header on a new file."""
    rows = list(rows)
    if not rows:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    new = not path.exists()
    with path.open("a", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=[*COLUMNS, *EXTRA])
        if new:
            writer.writeheader()
        writer.writerows(rows)


def read_csv(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise RecordError(f"no testing record at {path}: run the extraction command first")
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _number(value: Any) -> Any:
    """A CSV cell as the workbook should hold it: a number where it is one."""
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return value
    return int(number) if number.is_integer() and "." not in str(value) else number


def append_to_workbook(rows: Sequence[dict[str, Any]], xlsx: Path, *, sheet: str = SHEET) -> dict[str, int]:
    """Write ``rows`` into the workbook's sheet; returns how many were written and skipped.

    Columns are found by their header text, so a reordered sheet still fills
    right. A row lands in the first row whose data columns are empty and whose
    round is empty or already its own (the sheet's pre-typed "Baseline" rows),
    so nothing typed by hand is overwritten. A document already recorded for
    the same round and date is skipped: running the command twice adds nothing.
    """
    from copy import copy

    import openpyxl

    book = openpyxl.load_workbook(xlsx)
    # Matched without surrounding spaces, as the headers are: the team's sheet
    # is named "Testing Record " as typed.
    names = [name for name in book.sheetnames if name.strip() == sheet.strip()]
    if not names:
        raise RecordError(f"{xlsx} has no sheet {sheet!r} (it has {book.sheetnames})")
    ws = book[names[0]]
    headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value is not None}
    missing = [name for name in COLUMNS if name not in headers]
    if missing:
        raise RecordError(f"sheet {sheet!r} has no column {missing}; its headers are {sorted(headers)}")
    col = {name: headers[name] for name in COLUMNS}
    data = [name for name in COLUMNS if name != "Training Round"]

    def cell(row: int, name: str) -> Any:
        value = ws.cell(row=row, column=col[name]).value
        return value.strip() if isinstance(value, str) else value

    def key(round_label: Any, doc: Any, date: Any) -> tuple[str, str, str]:
        return (str(round_label or "").strip(), str(doc or "").strip(), str(date or "").strip()[:10])

    recorded = {key(cell(r, "Training Round"), cell(r, "Doc ID"), cell(r, "Training/Testing Date"))
                for r in range(2, ws.max_row + 1) if cell(r, "Doc ID")}
    style_row = 2
    written = skipped = 0
    for row in rows:
        if key(row["Training Round"], row["Doc ID"], row["Training/Testing Date"]) in recorded:
            skipped += 1
            continue
        target = next(r for r in range(2, ws.max_row + 2)
                      if all(cell(r, name) in (None, "") for name in data)
                      and cell(r, "Training Round") in (None, "", str(row["Training Round"]).strip()))
        for name in COLUMNS:
            out = ws.cell(row=target, column=col[name])
            if not out.has_style and ws.cell(row=style_row, column=col[name]).has_style:
                out._style = copy(ws.cell(row=style_row, column=col[name])._style)
            out.value = _number(row.get(name))
        recorded.add(key(row["Training Round"], row["Doc ID"], row["Training/Testing Date"]))
        written += 1
    book.save(xlsx)
    return {"written": written, "skipped": skipped}


#: Which bundles a staged run reads (``--select``). ``baseline``: every test and
#: validation document, and the training documents of every line those leave
#: under ~10 - the pilot baseline's documents per line (pilot/zero_shot_baseline.py).
SELECTIONS = ("baseline", "test", "test+val", "all")
BASELINE_PER_LINE = 10
INDEX_FILE = "doc_index.csv"
INDEX_COLUMNS = ("doc_id", "bundle", "split", "lob")


def doc_id(bundle: str, lob: str) -> str:
    """A short, stable id for a bundle: its line and a hash of its name. A
    bundle's name is its file name, which can carry an insured's name; the id
    goes into the shared workbook, the name stays in the local index."""
    from hashlib import sha256

    return f"{lob}_{sha256(bundle.encode('utf-8')).hexdigest()[:8]}"


def _bundles(bundles: Path) -> list[dict[str, Any]]:
    import json

    found = []
    for meta in sorted(bundles.glob("*/metadata.json")):
        folder = meta.parent
        if not (folder / "document.pdf").is_file() or not (folder / "golden.json").is_file():
            continue
        metadata = json.loads(meta.read_text(encoding="utf-8"))
        found.append({"bundle": folder.name, "folder": folder, "split": str(metadata.get("split") or ""),
                      "lob": str(metadata.get("lob") or "")})
    return found


def select_bundles(found: list[dict[str, Any]], select: str) -> list[dict[str, Any]]:
    if select not in SELECTIONS:
        raise RecordError(f"--select takes one of {SELECTIONS}, not {select!r}")
    if select == "all":
        return found
    held = [b for b in found if b["split"] in (("test",) if select == "test" else ("test", "val"))]
    if select != "baseline":
        return held
    per_line: dict[str, int] = {}
    for bundle in held:
        per_line[bundle["lob"]] = per_line.get(bundle["lob"], 0) + 1
    short = {line for line in {b["lob"] for b in found} if per_line.get(line, 0) < BASELINE_PER_LINE}
    return held + [b for b in found if b["split"] == "train" and b["lob"] in short]


def stage(bundles: Path, out: Path, *, select: str = "baseline") -> list[dict[str, Any]]:
    """Lay the chosen bundles out as ``latency_check`` reads them:
    ``<out>/pdfs/<id>.pdf`` with ``<out>/Gold_json/<id>.json`` (its gold lookup),
    ``<out>/docs.txt`` (``--doc`` and ``<pdf>=<line>`` per document, one argument
    a line: ``latency_check "@<out>/docs.txt"``) and
    ``<out>/doc_index.csv`` (id -> bundle, split, line). Run it again to restage."""
    import shutil

    chosen = select_bundles(_bundles(bundles), select)
    if not chosen:
        raise RecordError(f"no bundles in {bundles} for --select {select}")
    (out / "pdfs").mkdir(parents=True, exist_ok=True)
    (out / "Gold_json").mkdir(parents=True, exist_ok=True)
    index, docs = [], []
    for bundle in chosen:
        ident = doc_id(bundle["bundle"], bundle["lob"])
        pdf = out / "pdfs" / f"{ident}.pdf"
        shutil.copyfile(bundle["folder"] / "document.pdf", pdf)
        shutil.copyfile(bundle["folder"] / "golden.json", out / "Gold_json" / f"{ident}.json")
        index.append({"doc_id": ident, "bundle": bundle["bundle"], "split": bundle["split"], "lob": bundle["lob"]})
        docs += ["--doc", f"{pdf.resolve().as_posix()}={bundle['lob']}"]
    with (out / INDEX_FILE).open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(INDEX_COLUMNS))
        writer.writeheader()
        writer.writerows(index)
    (out / "docs.txt").write_text("\n".join(docs) + "\n", encoding="utf-8")
    return index


def rows_from_latency_check(run_dir: Path, index: Path, *, round_label: str, gold_dir: Path | None = None,
                            date: str | None = None) -> list[dict[str, Any]]:
    """The testing record of a ``testing.latency_check`` run: one row per document
    read (its ``<document>.json``), its split and line from the staged index and
    its carrier from the staged gold. A document that failed is left out and
    reported by ``latency_check`` itself."""
    import json

    with index.open(encoding="utf-8", newline="") as handle:
        known = {row["doc_id"]: row for row in csv.DictReader(handle)}
    gold_dir = gold_dir or index.parent / "Gold_json"
    rows = []
    for path in sorted(run_dir.glob("*.json")):
        if path.name in ("summary.json", "run_to_run.json"):
            continue
        report = json.loads(path.read_text(encoding="utf-8"))
        if "error" in report or "document" not in report:
            continue
        ident = Path(report["document"]).stem
        entry = known.get(ident, {})
        gold_path = gold_dir / f"{ident}.json"
        golden = json.loads(gold_path.read_text(encoding="utf-8")) if gold_path.is_file() else None
        model_s = (report.get("timings_s") or {}).get("model")
        result = SimpleNamespace(source_id=ident, latency_ms=model_s * 1000 if model_s is not None else None,
                                 window_stats=report.get("window_stats") or [], model_version=report.get("model"),
                                 mode=report.get("mode"))
        accuracy = report.get("accuracy") or {}
        rows.append(record_row(result, accuracy, metadata={"split": entry.get("split"), "lob": entry.get("lob")
                                                            or report.get("lob")},
                               golden=golden, pages=int(report.get("pages") or 0), round_label=round_label,
                               date=date or datetime.fromtimestamp(path.stat().st_mtime, UTC).date().isoformat()))
    return rows


def main(argv: Iterable[str] | None = None) -> int:  # pragma: no cover - thin CLI
    from testing.run_extraction import TESTING_ROOT, metrics_dir

    parser = argparse.ArgumentParser(description="Copy an extraction run's testing record into the workbook")
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--model", help="an extract run's model tag, as under testing/metrics/")
    source.add_argument("--latency-dir", dest="latency_dir", type=Path,
                        help="a testing.latency_check run's --out folder (no project code on the pod)")
    source.add_argument("--stage", action="store_true",
                        help="lay out --bundles for latency_check under --out (no workbook written)")
    parser.add_argument("--xlsx", type=Path, help="the team's training & test record workbook")
    parser.add_argument("--round", dest="round_label", default=None,
                        help="the Training Round to file the rows under (default: as recorded by the run)")
    parser.add_argument("--sheet", default=SHEET)
    parser.add_argument("--index", type=Path, help=f"with --latency-dir: the staged {INDEX_FILE}")
    parser.add_argument("--bundles", type=Path, help="with --stage: the bundles folder (data/bundles_c1)")
    parser.add_argument("--select", default="baseline", choices=SELECTIONS, help="with --stage: which bundles")
    parser.add_argument("--out", type=Path, help="with --stage: where to lay them out")
    args = parser.parse_args(list(argv) if argv is not None else None)

    if args.stage:
        if not args.bundles or not args.out:
            parser.error("--stage needs --bundles and --out")
        index = stage(args.bundles, args.out, select=args.select)
        by = {}
        for row in index:
            by[(row["lob"], row["split"])] = by.get((row["lob"], row["split"]), 0) + 1
        print(f"{len(index)} document(s) staged in {args.out}: "
              + ", ".join(f"{lob} {split} {n}" for (lob, split), n in sorted(by.items())))
        print(f"--doc arguments for latency_check: \"@{args.out / 'docs.txt'}\"")
        return 0
    if not args.xlsx:
        parser.error("--xlsx is required to write the workbook")
    if args.latency_dir:
        if not args.index or not args.round_label:
            parser.error("--latency-dir needs --index (the staged doc_index.csv) and --round")
        rows = rows_from_latency_check(args.latency_dir, args.index, round_label=args.round_label)
    else:
        rows = read_csv(metrics_dir(args.model, TESTING_ROOT) / RECORD_FILE)
    if args.round_label:
        rows = [{**row, "Training Round": args.round_label} for row in rows]
    counts = append_to_workbook(rows, args.xlsx, sheet=args.sheet)
    print(f"{counts['written']} row(s) written to {args.xlsx} [{args.sheet}], "
          f"{counts['skipped']} already there")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
